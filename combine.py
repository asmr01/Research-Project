import csv, collections
import enrich  # reuse code->name (SPEC), code->group (SPEC_GROUP), provider_type(code), county(), COUNTY_REGION
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

UP = "/root/.claude/uploads/87f376ba-3f62-5e40-825b-88b9f5b8ff90/"
DIR_CSV = UP + "825fafb9-optum_directory_ny_1.csv"        # Optum directory full roster (CareMount+ProHEALTH)
NPPES_CSV = UP + "c21cdcfb-optum_ny_providers.csv"         # NPPES (for Crystal Run)
LOOKUP_CSV = UP + "3d1017ff-nppes_lookup.csv"             # NPPES by-NPI lookup (tenure, secondary specialties)
OUT = "/home/user/Research-Project/Optum_NY_All_Providers.xlsx"

# --- directory: name-based specialty group + credential-based provider type ---
PC,MS,SS,WH,BH,AP,AH,OT="Primary Care","Medical Specialties","Surgical Specialties","Women's Health","Behavioral Health","Anesthesia/Pain/PM&R","Allied Health","Other"
CARD,ONC="Cardiology","Oncology"
def name_group(s):
    n=s.lower()
    if "cardio" in n or "cardiac" in n: return CARD
    if "oncology" in n: return ONC
    if any(k in n for k in ["orthopaedic","orthopedic","surgery","surgical","ophthalmology","otolaryngology","urology","retina"]): return SS
    if any(k in n for k in ["obstetric","gynecolog","midwife","maternal","pelvic"]): return WH
    if any(k in n for k in ["psycholog","behavioral health","counsel","psychiatr"]): return BH
    if any(k in n for k in ["pain medicine","physical medicine","rehabilitation","anesthe"]): return AP
    if any(k in n for k in ["physical therapist","podiatr","audiolog","optometr","nutrition","dietitian","therapist"]): return AH
    if any(k in n for k in ["family medicine","internal medicine","pediatrics","primary care","geriatric","obesity","developmental"]): return PC
    return MS
def cred_ptype(name, specialty):
    cred = name.split(",",1)[1].upper() if "," in name else ""
    toks = set(t.strip() for t in cred.replace("-"," ").split())
    if {"MD","DO"} & toks: return "Physician (MD/DO)"
    if {"NP","FNP","DNP","APRN","ANP","PNP","AGNP","WHNP"} & toks: return "Nurse Practitioner"
    if {"PA","PAC","RPA"} & toks: return "Physician Assistant"
    if "DPM" in toks: return "Podiatrist (DPM)"
    if {"DDS","DMD"} & toks: return "Dentist"
    if {"DPT","PT"} & toks: return "Therapy & Rehab"
    if {"OD","AUD","CNM","CM","RD","RDN"} & toks: return "Other Clinical"
    if {"PHD","PSYD"} & toks or "psycholog" in specialty.lower(): return "Behavioral Health (non-MD)"
    sl=specialty.lower()
    if "midwife" in sl: return "Other Clinical"
    if "physical therapist" in sl: return "Therapy & Rehab"
    if any(k in sl for k in ["audiolog","optometr","nutrition"]): return "Other Clinical"
    return "Physician (MD/DO)"
def parse_city(addr):
    parts=[p.strip() for p in addr.split(",")]
    return parts[-2] if len(parts)>=2 else ""
# extend enrich's county map with towns only present in the directory data
for _city,_cty in {"pomona":"Rockland","hewlett":"Nassau","north merrick":"Nassau","cedarhurst":"Nassau",
        "port jefferson station":"Suffolk","port jefferson":"Suffolk","miller place":"Suffolk",
        "centereach":"Suffolk","wading river":"Suffolk","lake katrine":"Ulster","seaford":"Nassau",
        "croton on hudson":"Westchester","briarcliff manor":"Westchester"}.items():
    enrich.CITY_COUNTY.setdefault(_city,_cty)

COLS = ["npi","display_name","first_name","last_name","gender","provider_type","specialty",
        "specialty_group","secondary_specialties","enumeration_date","years_since_npi",
        "accepting_new_patients","employed_or_contract","average_rating",
        "review_count","languages","cdo","region","county","city","zip","address","phone",
        "primary_location_name","num_locations","website_url","schedule_url","source"]

merged = {}  # npi -> row

# 1) Optum directory rows (authoritative, rich) -- CareMount + ProHEALTH
for r in csv.DictReader(open(DIR_CSV, encoding="utf-8-sig")):
    npi=r["npi"].strip()
    if not npi: continue
    city=parse_city(r["address"]); cty=enrich.county(city, r["zip"])
    merged[npi]={"npi":npi,"display_name":r["display_name"],"first_name":r["first_name"],
        "last_name":r["last_name"],"gender":r["gender"],
        "provider_type":cred_ptype(r["display_name"], r["specialty"]),"specialty":r["specialty"],
        "specialty_group":name_group(r["specialty"]),"accepting_new_patients":r["accepting_new_patients"],
        "employed_or_contract":r["employed_or_contract"],"average_rating":r["average_rating"],
        "review_count":r["review_count"],"languages":r["languages"],"cdo":r["cdo"],
        "region":enrich.COUNTY_REGION.get(cty,""),"county":cty,"city":city,"zip":r["zip"],
        "address":r["address"],"phone":r["phone"],"primary_location_name":r["primary_location_name"],
        "num_locations":r["num_locations"],"website_url":r["website_url"],"schedule_url":r["schedule_url"],
        "source":"Optum directory"}

# 2) Crystal Run from NPPES (specialty is a code) -- only add NPIs not already present
cr_added=0
for r in csv.DictReader(open(NPPES_CSV, encoding="utf-8-sig")):
    if "crystal run" not in r["brand"].lower(): continue
    npi=r["provider_npi"].strip()
    if not npi or npi in merged: continue
    code=r["specialty"]; cty=enrich.county(r["city"], r["zip"])
    nm=f'{r["first_name"]} {r["last_name"]}'.strip()
    if r["credential"].strip(): nm += ", " + r["credential"].strip()
    merged[npi]={"npi":npi,"display_name":nm,"first_name":r["first_name"],"last_name":r["last_name"],
        "gender":"","provider_type":enrich.provider_type(code),"specialty":enrich.SPEC.get(code,code),
        "specialty_group":enrich.SPEC_GROUP.get(code,"Other"),"accepting_new_patients":"",
        "employed_or_contract":"","average_rating":"","review_count":"","languages":"",
        "cdo":"Crystal Run","region":enrich.COUNTY_REGION.get(cty,""),"county":cty,"city":r["city"],
        "zip":r["zip"],"address":f'{r["street"]}, {r["city"]}, NY {r["zip"]}',"phone":r["phone"],
        "primary_location_name":r["location_name"],"num_locations":"","website_url":"","schedule_url":"",
        "source":"NPPES (Crystal Run)"}
    cr_added+=1

# 3) attach NPPES by-NPI lookup (tenure, sex fill, secondary specialties)
lk = {r["npi"].strip(): r for r in csv.DictReader(open(LOOKUP_CSV, encoding="utf-8-sig"))}
matched = 0
for npi, row in merged.items():
    L = lk.get(npi)
    if L:
        matched += 1
        row["enumeration_date"] = L["enumeration_date"]
        y = L["enumeration_date"][-4:]
        row["years_since_npi"] = (2026 - int(y)) if y.isdigit() else ""
        row["secondary_specialties"] = "; ".join(enrich.SPEC.get(c, c) for c in L["secondary_taxonomies"].split(";") if c)
        if not row.get("gender"):
            row["gender"] = {"M": "Male", "F": "Female"}.get(L["sex"], "")
    else:
        row["enumeration_date"] = ""; row["years_since_npi"] = ""; row["secondary_specialties"] = ""

# --- write flat workbook (no pivots) ---
wb=Workbook(); ws=wb.active; ws.title="Providers"
ws.append(COLS)
for row in sorted(merged.values(), key=lambda x:(x["county"], x["specialty_group"], x["last_name"], x["first_name"])):
    ws.append([row[c] for c in COLS])
for c in ws[1]: c.font=Font(bold=True); c.alignment=Alignment(vertical="center")
ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
for i in range(1, ws.max_column+1):
    w=max((len(str(ws.cell(row=r,column=i).value or "")) for r in range(1,min(ws.max_row,400)+1)),default=10)
    ws.column_dimensions[get_column_letter(i)].width=min(max(w+2,11),46)
wb.save(OUT)

print("Saved",OUT)
print("total providers:",len(merged),"| directory:",len(merged)-cr_added,"| Crystal Run added:",cr_added)
print("source:",dict(collections.Counter(v["source"] for v in merged.values())))
print("specialty_group:",dict(collections.Counter(v["specialty_group"] for v in merged.values())))
print("county:",dict(collections.Counter(v["county"] for v in merged.values())))
print("unmapped county:",{(v["city"],v["zip"]) for v in merged.values() if not v["county"]})
print("Cardiology:",sum(v["specialty_group"]=="Cardiology" for v in merged.values()),
      "| Oncology:",sum(v["specialty_group"]=="Oncology" for v in merged.values()))
print("NPPES lookup matched:",matched,"| with secondary specialties:",
      sum(1 for v in merged.values() if v["secondary_specialties"]),
      "| gender filled:",sum(1 for v in merged.values() if v["gender"]))
