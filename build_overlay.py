import collections
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

SRC="/home/user/Research-Project/Optum_NY_All_Providers.xlsx"
OUT="/home/user/Research-Project/Optum_NY_Comparison.xlsx"

P=load_workbook(SRC)["Providers"]; h=[c.value for c in P[1]]
ci=h.index("county"); pti=h.index("provider_type")
rows=list(P.iter_rows(min_row=2,values_only=True))
OUR=collections.Counter((r[pti],r[ci]) for r in rows)
our_cty=collections.Counter(r[ci] for r in rows)

PRISM={("Bronx","Physician (MD/DO)"):1,("Bronx","Nurse Practitioner"):1,("Bronx","Physician Assistant"):1,
 ("Kings","Physician (MD/DO)"):6,
 ("Nassau","Physician (MD/DO)"):70,("Nassau","Nurse Practitioner"):3,("Nassau","Physician Assistant"):10,("Nassau","Therapy & Rehab"):8,("Nassau","Other Clinical"):1,("Nassau","APC (unspecified)"):10,
 ("New York","Physician (MD/DO)"):28,("New York","Nurse Practitioner"):3,("New York","Physician Assistant"):4,("New York","Therapy & Rehab"):24,
 ("Queens","Physician (MD/DO)"):10,("Queens","Nurse Practitioner"):1,("Queens","Physician Assistant"):3,("Queens","Therapy & Rehab"):4,
 ("Rockland","Physician (MD/DO)"):3,
 ("Suffolk","Physician (MD/DO)"):13,("Suffolk","Nurse Practitioner"):6,("Suffolk","Physician Assistant"):10,("Suffolk","APC (unspecified)"):1,
 ("Westchester","Physician (MD/DO)"):27,("Westchester","Nurse Practitioner"):3,("Westchester","Physician Assistant"):3,("Westchester","Therapy & Rehab"):1,("Westchester","APC (unspecified)"):4}
prism_counties={c for (c,p) in PRISM}

MD="Physician (MD/DO)"
APC_SUB=["Nurse Practitioner","Physician Assistant","Therapy & Rehab","Other Clinical",
         "Podiatrist (DPM)","Behavioral Health (non-MD)","Other / Non-clinical","Dentist"]
DISPLAY={"New York":"New York (Manhattan)","Kings":"Kings (Brooklyn)","Richmond":"Richmond (Staten Island)"}
def disp(c): return DISPLAY.get(c,c)
prism_sorted=sorted(prism_counties,key=lambda c:-our_cty[c])
nonprism_sorted=sorted([c for c in our_cty if c not in prism_counties],key=lambda c:-our_cty[c])
COUNTIES=prism_sorted+nonprism_sorted

def ours(pt,c): return OUR.get((pt,c),0)
def prism(pt,c): return PRISM.get((c,pt),0)
def ours_apc(c): return sum(ours(pt,c) for pt in APC_SUB)
def prism_apc(c): return sum(prism(pt,c) for pt in APC_SUB)+prism("APC (unspecified)",c)

GREY=PatternFill("solid",fgColor="D9D9D9"); PINK=PatternFill("solid",fgColor="FCE4EC")
RED=PatternFill("solid",fgColor="F4B6C2"); HDR=PatternFill("solid",fgColor="DDEBF7")
TOT=PatternFill("solid",fgColor="E2EFDA"); YEL=PatternFill("solid",fgColor="FFF2CC")
thin=Side(style="thin",color="BFBFBF"); BORD=Border(thin,thin,thin,thin)
bold=Font(bold=True); ctr=Alignment(horizontal="center",vertical="center"); pct="0.0%"

wb=Workbook()

# ===== Sheet 1: Comparison (visual P / ours) =====
ws=wb.active; ws.title="Comparison"
ws["A1"]="Optum NY — Provider Type × County:  Prism / our roster"; ws["A1"].font=Font(bold=True,size=13,color="1F4E78")
ws["A2"]=("Cell = Prism / our-roster.  PINK = overlap.  RED = Prism > ours (resolve).  "
          "GREY cols = county not in Prism (our count only).  APC-unspecified = Prism untitled (compare at All-APC level).")
ws["A2"].font=Font(italic=True,color="C0006C")
ROWS=[MD]+APC_SUB+["APC – unspecified (Prism)","All APC (non-MD)","Grand Total"]
r0=4
ws.cell(r0,1,"Provider Type \\ County").font=bold
for j,c in enumerate(COUNTIES,2):
    cc=ws.cell(r0,j,disp(c)); cc.font=bold; cc.alignment=ctr; cc.border=BORD
    cc.fill=HDR if c in prism_counties else GREY
ws.cell(r0,len(COUNTIES)+2,"Total").font=bold; ws.cell(r0,len(COUNTIES)+2).fill=HDR; ws.cell(r0,len(COUNTIES)+2).alignment=ctr; ws.cell(r0,len(COUNTIES)+2).border=BORD
for i,rk in enumerate(ROWS,r0+1):
    ws.cell(i,1,rk).font=bold; ws.cell(i,1).border=BORD
    ot=pt_=0
    for j,c in enumerate(COUNTIES,2):
        if rk=="All APC (non-MD)": o,pr=ours_apc(c),prism_apc(c)
        elif rk=="Grand Total": o,pr=our_cty[c],sum(prism(x,c) for x in [MD]+APC_SUB)+prism("APC (unspecified)",c)
        elif rk=="APC – unspecified (Prism)": o,pr=0,prism("APC (unspecified)",c)
        else: o,pr=ours(rk,c),prism(rk,c)
        ot+=o; pt_+=pr
        cell=ws.cell(i,j); cell.alignment=ctr; cell.border=BORD
        if c not in prism_counties:
            cell.value=(o or ""); cell.fill=GREY
        elif o==0 and pr==0:
            cell.value=""; cell.fill=GREY
        else:
            cell.value=f"{pr} / {o}"
            if rk in (MD,"All APC (non-MD)") and pr>o: cell.fill=RED
            elif rk=="APC – unspecified (Prism)": cell.fill=YEL
            else: cell.fill=PINK
    e=ws.cell(i,len(COUNTIES)+2,f"{pt_} / {ot}"); e.font=bold; e.fill=TOT; e.alignment=ctr; e.border=BORD
    if rk in ("All APC (non-MD)","Grand Total"):
        for cc in ws[i]: cc.font=bold
ws.column_dimensions["A"].width=26
for j in range(2,len(COUNTIES)+3): ws.column_dimensions[get_column_letter(j)].width=12
ws.freeze_panes="B5"

# ===== numeric sheet builder (with formula totals + %) =====
def numeric_sheet(name, getter, title):
    s=wb.create_sheet(name)
    s["A1"]=title; s["A1"].font=Font(bold=True,size=12,color="1F4E78")
    hr=3; s.cell(hr,1,"Provider Type \\ County").font=bold
    pts=[MD]+APC_SUB+["APC (unspecified)"]
    for j,c in enumerate(COUNTIES,2):
        cc=s.cell(hr,j,disp(c)); cc.font=bold; cc.fill=HDR; cc.alignment=ctr; cc.border=BORD
    totcol=len(COUNTIES)+2; pctcol=totcol+1
    s.cell(hr,totcol,"Total").font=bold; s.cell(hr,totcol).fill=HDR; s.cell(hr,totcol).border=BORD
    s.cell(hr,pctcol,"% of all").font=bold; s.cell(hr,pctcol).fill=HDR; s.cell(hr,pctcol).border=BORD
    first=hr+1
    for i,pt in enumerate(pts,first):
        s.cell(i,1,pt).font=bold; s.cell(i,1).border=BORD
        for j,c in enumerate(COUNTIES,2):
            cell=s.cell(i,j,getter(pt,c)); cell.alignment=ctr; cell.border=BORD
        L=get_column_letter(2); R=get_column_letter(len(COUNTIES)+1)
        s.cell(i,totcol).value=f"=SUM({L}{i}:{R}{i})"; s.cell(i,totcol).border=BORD; s.cell(i,totcol).font=bold
    last=first+len(pts)-1; trow=last+1
    grand=f"{get_column_letter(totcol)}{trow}"
    s.cell(trow,1,"Total").font=bold; s.cell(trow,1).fill=TOT; s.cell(trow,1).border=BORD
    for j in range(2,totcol+1):
        col=get_column_letter(j)
        s.cell(trow,j).value=f"=SUM({col}{first}:{col}{last})"; s.cell(trow,j).fill=TOT; s.cell(trow,j).font=bold; s.cell(trow,j).border=BORD
    # % of all column (each row total / grand)
    for i in range(first,last+1):
        cell=s.cell(i,pctcol,f"={get_column_letter(totcol)}{i}/${get_column_letter(totcol)}${trow}")
        cell.number_format=pct; cell.border=BORD
    s.cell(trow,pctcol,f"={get_column_letter(totcol)}{trow}/${get_column_letter(totcol)}${trow}").number_format=pct
    s.column_dimensions["A"].width=26
    for j in range(2,pctcol+1): s.column_dimensions[get_column_letter(j)].width=11
    s.freeze_panes="B4"
    return s, first, last, trow, totcol  # positions for cross-refs

so,ofirst,olast,otrow,ototc=numeric_sheet("Ours (counts)",ours,"Our roster — Provider Type × County (counts)")
sp,pfirst,plast,ptrow,ptotc=numeric_sheet("Prism (counts)",prism,"Prism — Provider Type × County (counts)")

# ===== Summary with % formulas (per county: totals, %s, MD comparison) =====
ss=wb.create_sheet("Summary"); ss["A1"]="Summary — counts & percentages (live formulas)"; ss["A1"].font=Font(bold=True,size=12,color="1F4E78")
hdrs=["County","Our total","Our % of total","Our MDs","Our % of all MDs","% MDs within county",
      "Prism total","Prism MDs","MD gap (Prism−Ours)"]
hr=3
for j,t in enumerate(hdrs,1):
    c=ss.cell(hr,j,t); c.font=bold; c.fill=HDR; c.alignment=Alignment(horizontal="center",wrap_text=True,vertical="center"); c.border=BORD
oTcol=get_column_letter(ototc); pTcol=get_column_letter(ptotc)
for i,county in enumerate(COUNTIES,hr+1):
    cidx=COUNTIES.index(county)+2; col=get_column_letter(cidx)
    ss.cell(i,1,disp(county)).border=BORD
    ss.cell(i,2).value=f"='Ours (counts)'!{col}{otrow}"           # our total
    ss.cell(i,3).value=f"=B{i}/$B${hr+len(COUNTIES)+1}"            # our % of total
    ss.cell(i,4).value=f"='Ours (counts)'!{col}{ofirst}"          # our MDs (Physician row=first)
    ss.cell(i,5).value=f"=D{i}/$D${hr+len(COUNTIES)+1}"           # our % of all MDs
    ss.cell(i,6).value=f"=IFERROR(D{i}/B{i},0)"                   # % MDs within county
    ss.cell(i,7).value=f"='Prism (counts)'!{col}{ptrow}"          # prism total
    ss.cell(i,8).value=f"='Prism (counts)'!{col}{pfirst}"         # prism MDs
    ss.cell(i,9).value=f"=H{i}-D{i}"                              # MD gap
    for j in (3,5,6): ss.cell(i,j).number_format=pct
    for j in range(1,10): ss.cell(i,j).border=BORD; ss.cell(i,j).alignment=ctr
    ss.cell(i,1).alignment=Alignment(horizontal="left")
trow=hr+len(COUNTIES)+1
ss.cell(trow,1,"Total").font=bold; ss.cell(trow,1).fill=TOT
for j,L in [(2,"B"),(4,"D"),(7,"G"),(8,"H"),(9,"I")]:
    ss.cell(trow,j).value=f"=SUM({L}{hr+1}:{L}{trow-1})"; ss.cell(trow,j).font=bold; ss.cell(trow,j).fill=TOT
ss.cell(trow,3).value=f"=B{trow}/$B${trow}"; ss.cell(trow,5).value=f"=D{trow}/$D${trow}"; ss.cell(trow,6).value=f"=IFERROR(D{trow}/B{trow},0)"
for j in (3,5,6): ss.cell(trow,j).number_format=pct
for j in range(1,10): ss.cell(trow,j).border=BORD
# KPI callouts
k=trow+2
ss.cell(k,1,"% of all providers who are MDs:").font=bold; ss.cell(k,2).value=f"=D{trow}/B{trow}"; ss.cell(k,2).number_format=pct
ss.cell(k+1,1,"% of all providers in Nassau:").font=bold
nrow=hr+1+COUNTIES.index("Nassau"); ss.cell(k+1,2).value=f"=B{nrow}/B{trow}"; ss.cell(k+1,2).number_format=pct
ss.cell(k+2,1,"% of all MDs in Nassau:").font=bold; ss.cell(k+2,2).value=f"=D{nrow}/D{trow}"; ss.cell(k+2,2).number_format=pct
for col,w in zip("ABCDEFGHI",[20,11,12,9,12,14,11,10,16]): ss.column_dimensions[col].width=w
ss.freeze_panes="B4"

# ===== Resolve (real gaps: MD and All-APC where Prism>ours) =====
rs=wb.create_sheet("Resolve (Prism > ours)")
rs.append(["Level","County","Prism","Ours","Gap"])
for c in rs[1]: c.font=bold; c.fill=HDR; c.border=BORD
for c in prism_sorted:
    o,pr=ours(MD,c),prism(MD,c)
    if pr>o: rs.append(["Physician (MD/DO)",disp(c),pr,o,pr-o])
for c in prism_sorted:
    o,pr=ours_apc(c),prism_apc(c)
    if pr>o: rs.append(["All APC (non-MD)",disp(c),pr,o,pr-o])
for col,w in zip("ABCDE",[20,20,8,8,8]): rs.column_dimensions[col].width=w
for row in rs.iter_rows(min_row=2):
    for cell in row: cell.border=BORD
    row[-1].fill=RED; row[-1].font=bold

wb.save(OUT)
print("Saved",OUT)
print("Prism counties:",prism_sorted)
print("Greyed (not in Prism):",nonprism_sorted)
