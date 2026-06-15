import collections
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl import load_workbook

SRC = "/home/user/Research-Project/Optum_NY_All_Providers.xlsx"
OUT = "/home/user/Research-Project/Optum_NY_Comparison.xlsx"

P = load_workbook(SRC)["Providers"]; h = [c.value for c in P[1]]
ci=h.index("county"); pti=h.index("provider_type"); ti=h.index("target_specialty")
rows = list(P.iter_rows(min_row=2, values_only=True))

PT_ORDER=["Physician (MD/DO)","Nurse Practitioner","Physician Assistant","Therapy & Rehab",
          "Other Clinical","Podiatrist (DPM)","Behavioral Health (non-MD)","Other / Non-clinical","Dentist"]
TARGET_ORDER=["Allergy","Behavioral","Cardiology","Endocrinology","ENT","Gastroenterology","Hospitalist",
              "IMFM","Infectious Disease","OBGYN","Oncology/Radiation Oncology","Ophthalmology","Orthopedics",
              "Pain Management","Pediatrics","Physical Medicine and Rehab","Podiatry","Radiology","Rheumatology",
              "Surgery","Urgent Care","Urology","Other / no target"]
cty_tot=collections.Counter(r[ci] for r in rows)
COUNTIES=[c for c,_ in cty_tot.most_common()]   # ordered by size

GREY=PatternFill("solid",fgColor="D9D9D9")
HDR=PatternFill("solid",fgColor="DDEBF7"); TOTF=PatternFill("solid",fgColor="F2F2F2")
PINK=PatternFill("solid",fgColor="FCE4EC")
B=Border(*(Side(style="thin",color="BFBFBF"),)*4)
bold=Font(bold=True); ctr=Alignment(horizontal="center")

def matrix(ws, row_label, row_keys, counter, col_keys, title, note):
    ws["A1"]=title; ws["A1"].font=Font(bold=True,size=13,color="1F4E78")
    ws["A2"]=note; ws["A2"].font=Font(italic=True,color="C0006C")
    r0=4
    ws.cell(r0,1,row_label).font=bold
    for j,c in enumerate(col_keys,2):
        cc=ws.cell(r0,j,c); cc.font=bold; cc.fill=HDR; cc.alignment=ctr; cc.border=B
    tc=ws.cell(r0,len(col_keys)+2,"Total"); tc.font=bold; tc.fill=HDR; tc.alignment=ctr; tc.border=B
    coltot=collections.Counter(); grand=0
    for i,rk in enumerate(row_keys,r0+1):
        ws.cell(i,1,rk).font=bold; ws.cell(i,1).border=B
        rt=0
        for j,c in enumerate(col_keys,2):
            v=counter.get((rk,c),0); rt+=v; coltot[c]+=v
            cell=ws.cell(i,j,v); cell.alignment=ctr; cell.border=B
            if v==0: cell.fill=GREY
        e=ws.cell(i,len(col_keys)+2,rt); e.font=bold; e.fill=TOTF; e.alignment=ctr; e.border=B
        grand+=rt
    ti_=r0+len(row_keys)+1
    ws.cell(ti_,1,"Total").font=bold; ws.cell(ti_,1).fill=TOTF; ws.cell(ti_,1).border=B
    for j,c in enumerate(col_keys,2):
        e=ws.cell(ti_,j,coltot[c]); e.font=bold; e.fill=TOTF; e.alignment=ctr; e.border=B
    g=ws.cell(ti_,len(col_keys)+2,grand); g.font=bold; g.fill=TOTF; g.alignment=ctr; g.border=B
    ws.column_dimensions["A"].width=28
    for j in range(2,len(col_keys)+3): ws.column_dimensions[get_column_letter(j)].width=11
    ws.freeze_panes="B5"

wb=Workbook()
ws1=wb.active; ws1.title="County x ProviderType"
m1=collections.Counter((r[pti],r[ci]) for r in rows)
matrix(ws1,"Provider Type \\ County",PT_ORDER,m1,COUNTIES,
       "Optum NY — Providers by Provider Type × County (our base data)",
       "Grey = no providers. (Pending: pink 'included in Prism / total' overlay once Prism numbers are provided.)")

ws2=wb.create_sheet("Specialty x County")
m2=collections.Counter((r[ti],r[ci]) for r in rows)
matrix(ws2,"Specialty (your taxonomy) \\ County",TARGET_ORDER,m2,COUNTIES,
       "Optum NY — Providers by Specialty (your 22-cat taxonomy) × County (our base data)",
       "Grey = no providers. Specialty aligned to your external list via target_specialty.")

wb.save(OUT)
print("Saved",OUT)
print("counties:",COUNTIES)
print("grand total:",len(rows))
